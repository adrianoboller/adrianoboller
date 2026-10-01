# Renderizador aproximado (sem LibreOffice Impress no conteiner): retangulos,
# imagens e texto com quebra em Liberation Sans -- mais larga que a Calibri,
# entao o transbordo medido aqui e conservador.
import sys, io
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont
P = Presentation('PhxSql-telas-na-sequencia-de-uso.pptx')
DPI = 60; EMU = 914400
W = int(P.slide_width / EMU * DPI); H = int(P.slide_height / EMU * DPI)
def fnt(pt, b=False):
    f = '/usr/share/fonts/truetype/liberation/LiberationSans-%s.ttf' % ('Bold' if b else 'Regular')
    return ImageFont.truetype(f, max(6, int(pt * DPI / 72)))
def px(v): return int(v / EMU * DPI)
problemas = []
alvos = [int(a) for a in sys.argv[1:]] or range(1, len(P.slides) + 1)
for n, s in enumerate(P.slides, 1):
    im = Image.new('RGB', (W, H), '#010418'); d = ImageDraw.Draw(im)
    for sh in s.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if x < 0 or y < 0 or x + w > W + 1 or y + h > H + 1:
            problemas.append((n, 'fora do slide', sh.name))
        if sh.shape_type == 13:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA').resize((max(1,w), max(1,h)))
            im.paste(pic, (x, y), pic); continue
        try:
            c = sh.fill.fore_color.rgb; d.rounded_rectangle([x, y, x + w, y + h], 6, fill='#' + str(c))
        except Exception: pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            cy = y; usado = 0
            for p in sh.text_frame.paragraphs:
                runs = p.runs
                if not runs: continue
                r = runs[0]; pt = (r.font.size.pt if r.font.size else 18); b = bool(r.font.bold)
                cor = '#' + str(r.font.color.rgb) if r.font.color and r.font.color.type else '#DDE2EB'
                f = fnt(pt, b); txt = ''.join(x.text for x in runs)
                linhas = []; atual = ''
                for pal in txt.split(' '):
                    t = (atual + ' ' + pal).strip()
                    if d.textlength(t, font=f) <= w or not atual: atual = t
                    else: linhas.append(atual); atual = pal
                linhas.append(atual)
                lh = int(pt * 1.2 * DPI / 72)
                for l in linhas:
                    d.text((x, cy), l, font=f, fill=cor); cy += lh
                    if d.textlength(l, font=f) > w + 2: problemas.append((n, 'palavra mais larga que a caixa', l[:30]))
            if cy - y > h + 2: problemas.append((n, 'texto transborda %.0f%%' % (100 * (cy - y) / max(h,1)), sh.text_frame.text[:40]))
    if n in alvos: im.save('qa-%03d.png' % n)
print(len(P.slides), 'slides'); print('\n'.join(map(str, problemas)) or 'sem problemas')
